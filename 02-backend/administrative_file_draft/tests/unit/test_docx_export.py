"""Tests for the template-based DOCX export."""

import io
import zipfile
from types import SimpleNamespace
from uuid import uuid4

from defusedxml import ElementTree as DefusedElementTree
from docx import Document

from src.docx_export.generator import TEMPLATE_PATH, generate_vae_docx
from src.docx_export.template_utils import VAE_CONTENT_ANCHOR, WORDPROCESSINGML_NS


def _make_node(
    title: str,
    *,
    text: str = "",
    parent_id=None,
    order: float = 0,
):
    return SimpleNamespace(
        id=uuid4(),
        title=title,
        text=text,
        order=order,
        parent_vae_node_id=parent_id,
        is_hidden=False,
        is_deleted=False,
    )


def test_template_contains_placeholder_contract():
    """The copied DOCX template should expose the placeholder-based contract."""
    doc = Document(str(TEMPLATE_PATH))
    with zipfile.ZipFile(str(TEMPLATE_PATH)) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")
    table_text = "\n".join(paragraph.text for paragraph in doc.tables[0].rows[0].cells[0].paragraphs)

    # Placeholders replaced at export time (static values + PLS data)
    assert "[dokument_titel]" in document_xml
    assert "[vorhaben_traeger]" in document_xml
    assert "[vorhaben_name]" in document_xml
    assert "[Aktenzeichen]" in document_xml

    # Renamed placeholders (user fills in Word)
    assert "[Fernziel 1 \u2013 Fernziel 2]" in table_text
    assert "[Streckenabschnitt]" in table_text
    assert "[Gemarkungen]" in document_xml
    assert "[Ort ,Datum]" in document_xml

    # Unchanged placeholders
    assert "[verkehrseinheit]" in document_xml
    assert "[landkreis]" in document_xml

    assert VAE_CONTENT_ANCHOR in document_xml
    assert "{{ document_title }}" not in document_xml
    assert "{{ project_name }}" not in document_xml
    assert "{{ subtitle }}" not in document_xml


def test_generate_vae_docx_inserts_tree_removes_anchor_and_sets_update_fields():
    root = _make_node("A Root", text="Root body", order=0)
    child = _make_node("A.1 Child", text="Child body", parent_id=root.id, order=0)

    result = generate_vae_docx(
        [root, child],
        project_name="Neubau der BAB 123",
        internal_project_number="AZ-2026-042",
    )

    doc = Document(io.BytesIO(result))
    paragraph_texts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    joined_text = "\n".join(paragraph_texts)

    assert "A Root" in joined_text
    assert "Root body" in joined_text
    assert "A.1 Child" in joined_text
    assert "Child body" in joined_text
    assert VAE_CONTENT_ANCHOR not in joined_text
    assert paragraph_texts[-4:] == ["A Root", "Root body", "A.1 Child", "Child body"]

    with zipfile.ZipFile(io.BytesIO(result)) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")
        settings_xml = zf.read("word/settings.xml")

    assert VAE_CONTENT_ANCHOR not in document_xml
    assert "Planfeststellungsbeschluss" in document_xml
    assert "Neubau der BAB 123" in document_xml
    assert "AZ-2026-042" in document_xml
    assert "Die Autobahn GmbH des Bundes" in document_xml
    assert "[dokument_titel]" not in document_xml
    assert "[vorhaben_name]" not in document_xml
    assert "[Aktenzeichen]" not in document_xml
    assert "[vorhaben_traeger]" not in document_xml

    settings_root = DefusedElementTree.fromstring(settings_xml)
    update_fields = settings_root.find(f"{{{WORDPROCESSINGML_NS}}}updateFields")

    assert update_fields is not None
    assert update_fields.get(f"{{{WORDPROCESSINGML_NS}}}val") == "true"
