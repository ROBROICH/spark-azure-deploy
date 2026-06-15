"""Template-based DOCX generator for VAE node exports."""

import io
import logging
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph  # type: ignore[attr-defined]

from src.docx_export.template_utils import (
    VAE_CONTENT_ANCHOR,
    clear_content_from_anchor,
    enable_update_fields,
    find_anchor_paragraph,
    replace_placeholders_in_docx_xml,
)
from src.docx_export.vae_tree_builder import VaeTreeNode, build_tree
from src.models.db_models import VaeNode

logger = logging.getLogger(__name__)

TEMPLATE_PATH = Path(__file__).parent.parent / "template.docx"

_MAX_HEADING_DEPTH = 9


_STATIC_REPLACEMENTS: dict[str, str] = {
    "dokument_titel": "Planfeststellungsbeschluss",
    "vorhaben_traeger": "Die Autobahn GmbH des Bundes",
}


def generate_vae_docx(
    nodes: list[VaeNode],
    project_name: str,
    internal_project_number: str,
) -> bytes:
    """Generate a DOCX document from a flat list of VaeNode DB rows.

    Steps:
    1. Load the Musterbeschluss template.
    2. Render the VAE tree at the explicit body anchor.
    3. Delete the anchor paragraph and all sample content after it.
    4. Patch settings.xml so Word updates the TOC on open.
    5. Replace cover-page placeholders with project data and static values.

    Returns raw DOCX bytes.
    """
    doc = Document(str(TEMPLATE_PATH))

    tree = build_tree(nodes)
    anchor_paragraph = find_anchor_paragraph(doc, VAE_CONTENT_ANCHOR)
    _render_tree(anchor_paragraph, tree)
    clear_content_from_anchor(anchor_paragraph)

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = enable_update_fields(buf.getvalue())

    replacements = {
        **_STATIC_REPLACEMENTS,
        "vorhaben_name": project_name,
        "Aktenzeichen": internal_project_number,
    }
    return replace_placeholders_in_docx_xml(docx_bytes, replacements)


def _render_tree(anchor_paragraph: Paragraph, roots: list[VaeTreeNode]) -> None:
    for node in roots:
        _render_node(anchor_paragraph, node, depth=1)


def _render_node(anchor_paragraph: Paragraph, node: VaeTreeNode, depth: int) -> None:
    heading_depth = min(depth, _MAX_HEADING_DEPTH)
    anchor_paragraph.insert_paragraph_before(node.title, style=f"Heading {heading_depth}")

    if node.text.strip():
        for line in node.text.splitlines():
            stripped = line.strip()
            if stripped:
                anchor_paragraph.insert_paragraph_before(stripped)

    for child in node.children:
        _render_node(anchor_paragraph, child, depth + 1)
