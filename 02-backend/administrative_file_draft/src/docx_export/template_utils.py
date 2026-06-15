"""Low-level DOCX manipulation helpers for template-based generation."""

import io
import logging
import zipfile

from defusedxml import ElementTree as DefusedElementTree
from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph  # type: ignore[attr-defined]

logger = logging.getLogger(__name__)

WORDPROCESSINGML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
VAE_CONTENT_ANCHOR = "{{ VAE_CONTENT_START }}"


def find_anchor_paragraph(doc: DocxDocument, anchor_text: str = VAE_CONTENT_ANCHOR) -> Paragraph:
    """Return the paragraph that contains the body insertion anchor."""
    for para in doc.paragraphs:
        if para.text.strip() == anchor_text:
            return para

    raise RuntimeError(
        f"Anchor paragraph {anchor_text!r} not found in DOCX template. The template contract may have changed."
    )


def clear_content_from_anchor(anchor_paragraph: Paragraph) -> None:
    """Delete the anchor paragraph and every following body element.

    Preserves cover, TOC, and section properties (w:sectPr).
    Operates at the XML level because python-docx has no paragraph.delete().
    """
    anchor_element = anchor_paragraph._p
    body = anchor_element.getparent()
    removing = False
    to_remove = []
    for child in body:
        if child is anchor_element:
            removing = True
        if removing:
            local = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if local != "sectPr":
                to_remove.append(child)

    for elem in to_remove:
        body.remove(elem)


def enable_update_fields(docx_bytes: bytes) -> bytes:
    """Patch word/settings.xml so Word refreshes the TOC on open."""
    buf = io.BytesIO(docx_bytes)
    files: dict[str, bytes] = {}
    with zipfile.ZipFile(buf, "r") as zin:
        for name in zin.namelist():
            files[name] = zin.read(name)

    settings = files.get("word/settings.xml", b"")
    if not settings:
        logger.warning("word/settings.xml missing from DOCX — TOC auto-update not applied")
        return docx_bytes

    root = DefusedElementTree.fromstring(settings)
    update_fields = root.find(f"{{{WORDPROCESSINGML_NS}}}updateFields")

    if update_fields is None:
        update_fields = DefusedElementTree.fromstring(f'<w:updateFields xmlns:w="{WORDPROCESSINGML_NS}" />')
        root.append(update_fields)

    update_fields.set(f"{{{WORDPROCESSINGML_NS}}}val", "true")
    files["word/settings.xml"] = DefusedElementTree.tostring(root, encoding="utf-8", xml_declaration=True)

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    return out.getvalue()


def replace_placeholders_in_docx_xml(docx_bytes: bytes, replacements: dict[str, str]) -> bytes:
    """Replace placeholders directly in word/document.xml.

    This covers content python-docx does not expose well, such as text inside
    floating text boxes / drawing text containers on the cover page.
    """
    buf = io.BytesIO(docx_bytes)
    files: dict[str, bytes] = {}
    with zipfile.ZipFile(buf, "r") as zin:
        for name in zin.namelist():
            files[name] = zin.read(name)

    document_xml = files.get("word/document.xml", b"")
    if not document_xml:
        logger.warning("word/document.xml missing from DOCX — XML placeholder replacement skipped")
        return docx_bytes

    document_text = document_xml.decode("utf-8")
    for key, value in replacements.items():
        document_text = document_text.replace(f"[{key}]", value)

    files["word/document.xml"] = document_text.encode("utf-8")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    return out.getvalue()
